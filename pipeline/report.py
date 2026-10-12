import re

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.deseq2 import comparison_name

UNINFORMATIVE_DESCRIPTION = "uncharacterized|hypothetical"

COLUMNS = [
    "gene_id", "best_annotation", "annotation_source", "log2FoldChange", "padj", "baseMean",
    "gene_name", "gff_description", "gene_biotype", "unknown_in_gff",
    "blast_status", "blast_top_hit", "blast_name_support", "blast_best_single_hit",
]


def blast_status_for(gene_id, row, blast_summary, genes_to_blast, too_long):
    if not row["unknown_in_gff"]:
        return "not needed (named gene)"
    if gene_id in blast_summary.index:
        return blast_summary.at[gene_id, "blast_status"]
    if gene_id in too_long:
        return "not BLASTed (too long)"
    if gene_id in genes_to_blast:
        return "waiting for BLAST results"
    return "not BLASTed (biotype excluded)"


def write_final_tables(settings, significant, gene_table, blast_summary, genes_to_blast, too_long, final_dir):
    final_dir.mkdir(parents=True, exist_ok=True)
    genes_to_blast, too_long = set(genes_to_blast), set(too_long)
    counts = {}
    tables = {}

    for contrast, sig in significant.items():
        table = sig[["gene_id", "baseMean", "log2FoldChange", "padj"]].join(gene_table, on="gene_id")
        # Genes missing from the GFF count as unknown
        table["unknown_in_gff"] = table["unknown_in_gff"].ne(False)

        table["blast_status"] = [blast_status_for(g, row, blast_summary, genes_to_blast, too_long)
                                 for g, row in zip(table["gene_id"], table.to_dict("records"))]
        table["blast_top_hit"] = table["gene_id"].map(blast_summary["blast_top_hit"]).where(
            table["blast_status"] == "identified")
        identified = table["blast_status"] == "identified"
        table["blast_name_support"] = table["gene_id"].map(blast_summary["name_support"]).where(identified)
        table["blast_best_single_hit"] = table["gene_id"].map(blast_summary["best_single_hit"]).where(identified)

        # An informative GFF description wins; otherwise (none, or e.g. "uncharacterized
        # protein LOC123") the BLAST name is used, falling back to the GFF description
        has_desc = table["gff_description"].notna()
        informative_desc = has_desc & ~table["gff_description"].str.contains(
            UNINFORMATIVE_DESCRIPTION, case=False, na=False)
        from_blast = ~informative_desc & table["blast_top_hit"].notna()

        table["best_annotation"] = table["gff_description"].where(informative_desc, table["blast_top_hit"])
        table["best_annotation"] = table["best_annotation"].where(table["best_annotation"].notna(),
                                                                  table["gff_description"])
        table["annotation_source"] = ""
        table.loc[has_desc, "annotation_source"] = "GFF (uninformative)"
        table.loc[informative_desc & ~table["unknown_in_gff"], "annotation_source"] = "GFF"
        table.loc[informative_desc & table["unknown_in_gff"], "annotation_source"] = "GFF (description only)"
        table.loc[from_blast, "annotation_source"] = "BLAST"

        table = table.sort_values("padj")[COLUMNS]
        table.to_csv(final_dir / f"annotated_{comparison_name(contrast)}.csv", index=False)
        tables[comparison_name(contrast)] = table

        counts[comparison_name(contrast)] = {
            "significant": len(table),
            "unknown in GFF": int(table["unknown_in_gff"].sum()),
            "identified by BLAST": int((table["blast_status"] == "identified").sum()),
            "still unknown after BLAST": int((table["blast_status"] == "unknown").sum()),
            "waiting for BLAST": int((table["blast_status"] == "waiting for BLAST results").sum()),
            "no annotation": int(table["best_annotation"].isna().sum()),
        }

    write_excel(final_dir / "annotated_results.xlsx", tables, counts, settings)
    return counts


COLUMN_GUIDE = {
    "gene_id": "Gene ID from the GFF / count tables",
    "best_annotation": "GFF description if it is informative, otherwise the BLAST name",
    "annotation_source": "Where best_annotation came from",
    "log2FoldChange": "DESeq2 log2 fold change (positive = higher in the first condition)",
    "padj": "DESeq2 adjusted p-value",
    "baseMean": "DESeq2 mean normalized count",
    "gene_name": "GFF Name attribute",
    "gff_description": "GFF description attribute",
    "gene_biotype": "GFF gene_biotype attribute",
    "unknown_in_gff": "TRUE if the gene has no real name in the GFF",
    "blast_status": "identified / unknown / waiting for BLAST results / not needed / not BLASTed",
    "blast_top_hit": "Gene name chosen by the e-value-weighted vote of BLAST hits",
    "blast_name_support": "Share of the vote the name got (0-100%). Highlighted when under 50%",
    "blast_best_single_hit": "The single best (lowest e-value) BLAST hit, for comparison",
}
NUMBER_FORMATS = {"log2FoldChange": "0.00", "padj": "0.00E+00", "baseMean": "#,##0.0",
                  "blast_name_support": "0%"}
COLUMN_WIDTHS = {"gene_id": 18, "best_annotation": 45, "annotation_source": 20, "gff_description": 40,
                 "blast_status": 24, "blast_top_hit": 40, "blast_best_single_hit": 40}
HEADER_FILL = PatternFill("solid", start_color="DDE4EE")
LOW_SUPPORT_FILL = PatternFill("solid", start_color="FFE699")


def sheet_name(name, used):
    # Excel sheet names: max 31 characters, no []:*?/\ and unique
    name = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
    base, n = name, 2
    while name.lower() in used:
        name = f"{base[:28]}_{n}"
        n += 1
    used.add(name.lower())
    return name


def write_excel(path, tables, counts, settings):
    # Same tables as the CSVs, one sheet per comparison, plus a summary sheet
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_rows = [{"comparison": name, **numbers} for name, numbers in counts.items()]
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False, startrow=2)
        summary = writer.sheets["Summary"]
        summary["A1"] = "Significant genes per comparison"
        summary["A1"].font = Font(bold=True, size=13)
        summary["A2"] = (f"padj < {settings['padj_cutoff']} and |log2FC| > {settings['log2fc_cutoff']}. "
                         "One sheet per comparison, sorted by padj.")

        guide_row = len(summary_rows) + 6
        summary.cell(row=guide_row, column=1, value="Columns").font = Font(bold=True, size=13)
        for i, (column, meaning) in enumerate(COLUMN_GUIDE.items(), start=guide_row + 1):
            summary.cell(row=i, column=1, value=column).font = Font(bold=True)
            summary.cell(row=i, column=2, value=meaning)
        for cell in summary[3]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        summary.column_dimensions["A"].width = 30
        for col in range(2, 8):
            summary.column_dimensions[get_column_letter(col)].width = 16

        used = {"summary"}
        for name, table in tables.items():
            sheet = sheet_name(name, used)
            table.to_excel(writer, sheet_name=sheet, index=False)
            format_table_sheet(writer.sheets[sheet], table)


def format_table_sheet(ws, table):
    n_rows = len(table) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{n_rows}"

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for i, column in enumerate(table.columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(column, 14)
        if column in NUMBER_FORMATS:
            for (cell,) in ws.iter_rows(min_row=2, max_row=n_rows, min_col=i, max_col=i):
                cell.number_format = NUMBER_FORMATS[column]
        if column == "blast_name_support":
            # ISNUMBER so blank cells (no BLAST name) aren't treated as 0
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{n_rows}",
                FormulaRule(formula=[f"AND(ISNUMBER({letter}2),{letter}2<0.5)"], fill=LOW_SUPPORT_FILL))
