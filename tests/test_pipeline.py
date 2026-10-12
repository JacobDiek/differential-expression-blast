"""End-to-end test: runs the whole pipeline on the synthetic test dataset, twice
(before and after "downloading" BLAST results). Needs R with DESeq2."""
import json
import shutil
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from run_pipeline import run

TEST_DATA = Path(__file__).resolve().parent / "data"


def deseq2_available():
    if shutil.which("Rscript") is None:
        return False
    result = subprocess.run(["Rscript", "-e", 'quit(status = !requireNamespace("DESeq2", quietly = TRUE))'],
                            capture_output=True)
    return result.returncode == 0


pytestmark = pytest.mark.skipif(not deseq2_available(), reason="needs R with DESeq2")


@pytest.fixture(scope="module")
def pipeline_run(tmp_path_factory):
    # Copy the test data so the outputs go to a temporary folder
    data = tmp_path_factory.mktemp("run") / "data"
    shutil.copytree(TEST_DATA, data, ignore=shutil.ignore_patterns("output"))
    return data, json.loads((data / "expected.json").read_text())


def test_first_run_stops_at_blast(pipeline_run, capsys):
    data, expected = pipeline_run
    run(data / "test_config.yaml")
    out = capsys.readouterr().out
    output = data / "output"

    significant = pd.read_csv(output / "deseq2" / "significant_genes_treated_vs_control.csv")
    assert set(significant["gene_id"]) == set(expected["significant_genes"])

    batch = output / "ncbi" / "blast_batches" / "batch_1"
    assert (batch / "genes.txt").read_text().split() == expected["genes_to_blast"]
    assert len(list((batch / "chunks").glob("*.fasta"))) == 3
    assert (batch / "INSTRUCTIONS.txt").exists()
    assert (output / "ncbi" / "genes_too_long.txt").read_text().split() == expected["too_long"]
    assert "ACTION NEEDED" in out

    table = pd.read_csv(output / "final_tables" / "annotated_treated_vs_control.csv")
    waiting = table[table["blast_status"] == "waiting for BLAST results"]["gene_id"]
    assert sorted(waiting) == sorted(expected["genes_to_blast"])
    named = table[~table["unknown_in_gff"]]
    assert len(named) == 14
    assert (named["blast_status"] == "not needed (named gene)").all()
    assert (named["annotation_source"] == "GFF").all()
    too_long = table.set_index("gene_id").loc[expected["too_long"], "blast_status"]
    assert (too_long == "not BLASTed (too long)").all()

    record = json.loads((output / "run_info" / "latest_run.json").read_text())
    assert record["status"] == "waiting for BLAST results"
    assert record["settings"]["star_count_column"] == 4
    assert record["deseq2_run_with"]["DESeq2"]


def test_second_run_names_genes_from_blast_results(pipeline_run, capsys):
    data, expected = pipeline_run
    output = data / "output"
    for xml in (data / "blast_results").glob("*.xml"):
        shutil.copy(xml, output / "ncbi" / "ncbi_files_large_alignment")

    run(data / "test_config.yaml")
    out = capsys.readouterr().out
    assert "DESeq2 results are up to date, skipping." in out
    assert "Pipeline complete." in out

    table = pd.read_csv(output / "final_tables" / "annotated_treated_vs_control.csv",
                        keep_default_na=False).set_index("gene_id")
    for gene, exp in expected["blast"].items():
        assert table.at[gene, "blast_status"] == exp["status"], gene
        assert (table.at[gene, "blast_top_hit"] or None) == exp["name"], gene
        if exp["status"] == "identified":
            assert table.at[gene, "best_annotation"] == exp["name"]
            assert table.at[gene, "annotation_source"] == "BLAST"

    missing = (output / "ncbi" / "blast_batches" / "batch_1" / "missing_results.txt").read_text().split()
    assert missing == expected["missing_from_blast"]
    assert len(list((output / "ncbi" / "blast_batches").iterdir())) == 1   # no new batch

    workbook = load_workbook(output / "final_tables" / "annotated_results.xlsx")
    assert workbook.sheetnames == ["Summary", "treated_vs_control"]
    assert workbook["treated_vs_control"].max_row == len(expected["significant_genes"]) + 1

    record = json.loads((output / "run_info" / "latest_run.json").read_text())
    assert record["status"] == "complete"
    assert record["blast"]["searches"][0]["database"] == "core_nt"
    assert len(list((output / "run_info").glob("run_*.json"))) == 2
